"""Ground truth beside the system's numbers, as one spreadsheet.

    python -m scripts.results_sheet <sweep_dir> <ground_truth.json> <out.xlsx>

Sheets: "Pass bars" (every bar: target, truth, system, pass), "Books" (each true book next to the
book the system matched to it, with dimension errors) and "Not measured" (what this ground truth does
not cover). Needs openpyxl.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from scripts.evaluate import _same_title, evaluate

BOLD = Font(bold=True)
PASS, FAIL = PatternFill("solid", fgColor="D9F2D9"), PatternFill("solid", fgColor="F8D7D7")


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = "A2"
    return ws


def build(sweep: Path, truth_path: Path, out: Path) -> None:
    packet = json.loads((sweep / "claim_packet.json").read_text(encoding="utf-8"))
    truth = json.loads(truth_path.read_text(encoding="utf-8"))
    results = evaluate(packet, truth)
    wb = Workbook()
    wb.remove(wb.active)

    bars = [
        ("Book count", "within 5%", results["book_count"]["truth"], results["book_count"]["system"],
         f'{results["book_count"]["error_pct"]}% off', results["book_count"]["pass"]),
        ("Titles", ">= 70% of legible right, <= 3% confidently wrong", f'{results["titles"]["legible_truth"]} legible',
         f'{results["titles"]["correct"]} right ({results["titles"]["correct_pct"]}%)',
         f'{results["titles"]["confidently_wrong"]} confidently wrong ({results["titles"]["confidently_wrong_pct"]}%)',
         results["titles"]["pass"]),
        ("Spine dimensions", "20 measured spines within 15%", f'{results["spine_dimensions"]["sample"]} spines',
         f'{results["spine_dimensions"]["within_15pct"]} within 15%',
         f'median error: height {results["spine_dimensions"]["median_height_err_pct"]}%, '
         f'thickness {results["spine_dimensions"]["median_thickness_err_pct"]}%', results["spine_dimensions"]["pass"]),
        ("Time to packet", "under 5 minutes", "", f'{results["time_to_packet_s"]["system"]} s', "",
         results["time_to_packet_s"]["pass"]),
        ("Prices", "within 25% on 15 books", "not measured", "", "see 'Not measured'", None),
        ("Non-book items", ">= 80% found", "not measured", f'{len(packet["items"])} found', "see 'Not measured'", None),
        ("Floor / wall area", "within 10% / 15%", "not measured", "", "see 'Not measured'", None),
    ]
    ws = _sheet(wb, "Pass bars", ["Pass bar", "Target", "Ground truth", "System", "Error", "Pass"],
                [list(b[:5]) + ["PASS" if b[5] else ("FAIL" if b[5] is False else "n/a")] for b in bars],
                [18, 44, 18, 26, 52, 8])
    for row in ws.iter_rows(min_row=2):
        row[5].fill = PASS if row[5].value == "PASS" else FAIL if row[5].value == "FAIL" else PatternFill()

    books = packet["books"]
    unmatched = list(books)
    rows = []
    for t in truth["books"]:
        match = next((b for b in unmatched if t.get("title") and _same_title(b.get("title", ""), t["title"])), None)
        if match:
            unmatched.remove(match)
        h, th = (match or {}).get("spine_height_cm"), (match or {}).get("spine_thickness_cm")
        rows.append([
            f'{t["unit"]} · Shelf {t["shelf"]} · {t["position"]}', t["orientation"], "yes" if t.get("legible") else "no",
            t.get("title") or "(blank spine)", t.get("author", ""), t["height_cm"], t["thickness_cm"],
            (match or {}).get("id", ""), (match or {}).get("title", ""), (match or {}).get("id_confidence", ""), h, th,
            round(100 * abs(h - t["height_cm"]) / t["height_cm"], 1) if h else "",
            round(100 * abs(th - t["thickness_cm"]) / t["thickness_cm"], 1) if th else "",
        ])
    for b in unmatched:  # system books with no true counterpart (duplicates, unreadable spines, misreads)
        rows.append(["", "", "", "", "", "", "", b["id"], b.get("title") or "(unidentified)", b.get("id_confidence", ""),
                     b.get("spine_height_cm"), b.get("spine_thickness_cm"), "", ""])
    _sheet(wb, "Books", ["True position", "Orientation", "Legible", "True title", "True author", "True height cm",
                         "True thickness cm", "System book", "System title", "Confidence", "System height cm",
                         "System thickness cm", "Height error %", "Thickness error %"],
           rows, [24, 11, 8, 32, 22, 13, 15, 12, 32, 11, 15, 17, 14, 16])

    _sheet(wb, "Not measured", ["What", "Why"], [
        ["Hand-checked prices (15 books)", "No real room or books: the library is synthetic. Every price in the "
         "packet carries its listing URL and date for checking."],
        ["Non-book item list", "The synthetic shelves hold no objects besides the units; items found are listed in the packet."],
        ["Room tape measurements", "No real room. The mock flow's room is an example 4.2 x 3.6 x 2.7 m from simulated "
         "AR points, not a measurement."],
    ], [32, 110])
    wb.save(out)


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
